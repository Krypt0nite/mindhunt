#!/usr/bin/env python3
"""Export the cricket pack to an Excel sheet people can correct and extend.

Run: python3 tools/export_cricket_xlsx.py server/packs/cricket-men-intl.v3.json out.xlsx
Re-import with tools/import_cricket_xlsx.py.
"""
import json
import sys

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

FONT = "Arial"
HEAD = PatternFill("solid", fgColor="1B2A4A")
EDIT = PatternFill("solid", fgColor="FFF8DC")   # cells to correct
LOCK = PatternFill("solid", fgColor="E6E9EE")   # do not edit
NEW = PatternFill("solid", fgColor="EAF5EE")    # blank rows for new players
thin = Side(style="thin", color="C9D1DC")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)

COLUMNS = [
    # header, width, list name (for dropdown) or None, strict?
    ("ID", 11, None, False),
    ("Name", 24, None, False),
    ("Surname", 16, None, False),
    ("Country", 15, "countries", False),
    ("Role", 15, "roles", True),
    ("Batting hand", 12, "hands", True),
    ("Bowling", 20, "bowling", True),
    ("Bowling arm", 12, "hands", True),
    ("ODI debut decade", 13, "decades", True),
    ("Captained country", 12, "yesno", True),
    ("ODI World Cup winner", 13, "yesno", True),
    ("Keep in game", 10, "yesno", True),
    ("Verified", 10, "yesno", True),
    ("Source / notes", 34, None, False),
    ("Image page URL", 30, None, False),
    ("Image licence", 14, None, False),
    ("Image credit", 20, None, False),
]

NOTES = {
    "ID": "Leave as is. Blank for new players; an ID is assigned on import.",
    "Surname": "Used for 'Does their surname start with…' questions. Spaces and hyphens are ignored.",
    "Country": "The international side they played for. Pick from the list, or type a new one.",
    "Role": "Main role. Keepers who mainly bat are Wicketkeeper.",
    "Bowling": "Batters and wicketkeepers are 'Not a regular bowler' even if they bowled occasionally.",
    "Bowling arm": "Only for Pace or Spin bowlers. Leave blank for 'Not a regular bowler'.",
    "ODI debut decade": "Decade of their first ODI for any country (e.g. first ODI 2008 = 2000). N/A if they never played an ODI.",
    "Captained country": "Yes if they captained their country in any international format, including as stand-in.",
    "ODI World Cup winner": "Yes if they were in a squad that won the men's 50-over World Cup.",
    "Keep in game": "Set to No to remove the player.",
    "Verified": "Set to Yes once you have checked the row against a reliable source.",
    "Source / notes": "Where you checked it (e.g. ESPNcricinfo profile link), or anything that needs a decision.",
    "Image page URL": "Optional. A Wikimedia Commons file page, e.g. https://commons.wikimedia.org/wiki/File:...",
    "Image licence": "Licence shown on that page, e.g. CC BY-SA 4.0. Must allow reuse.",
    "Image credit": "Author/photographer as the licence requires.",
}


def main(src, dst):
    pack = json.load(open(src))
    attrs = {a["key"]: a for a in pack["attributes"]}
    countries = list(attrs["country"]["values"]) + [c for c in ["Ireland", "Zimbabwe", "Netherlands", "Scotland"] if c not in attrs["country"]["values"]]
    lists = {
        "countries": countries,
        "roles": attrs["role"]["values"],
        "hands": ["Left", "Right"],
        "bowling": attrs["bowling"]["values"],
        "decades": ["N/A"] + [str(y) for y in range(1970, 2030, 10)],
        "yesno": ["Yes", "No"],
    }

    wb = Workbook()

    # --- How to fill ---
    h = wb.active
    h.title = "How to fill"
    h.column_dimensions["A"].width = 26
    h.column_dimensions["B"].width = 95
    rows = [
        ("MindHunt cricket pack", ""),
        ("", ""),
        ("What this is", f"Every player in the game ({len(pack['candidates'])} today) with the facts players can ask about. Answers in the game come straight from this sheet, so a wrong cell means a wrong answer."),
        ("What to do", "1. Check each row in the Players tab and fix anything wrong (yellow cells)."),
        ("", "2. Set Verified to Yes for rows you have checked, and put your source in Source / notes."),
        ("", "3. Add new players in the green rows at the bottom (leave ID blank)."),
        ("", "4. To remove a player, set Keep in game to No. Do not delete the row."),
        ("", "5. Optional: add a Wikimedia Commons image page, licence and credit for the results screen."),
        ("", "6. Send the file back. It is converted into the game data automatically."),
        ("Do not edit", "The grey ID column and the header row."),
        ("Dropdowns", "Most columns only accept listed values. Country also accepts a new country if you type one."),
        ("", ""),
        ("Example new player", "Name: Rishabh Pant | Surname: Pant | Country: India | Role: Wicketkeeper | Batting hand: Left | Bowling: Not a regular bowler | Bowling arm: (blank) | ODI debut decade: 2010 | Captained country: Yes | ODI World Cup winner: No | Keep in game: Yes | Verified: No"),
        ("", "(Example only: check every value before adding a real row.)"),
        ("", ""),
        ("Column", "Meaning"),
    ]
    for col, *_ in COLUMNS:
        rows.append((col, NOTES.get(col, "")))
    for r in rows:
        h.append(list(r))
    for row in h.iter_rows():
        for c in row:
            c.font = Font(name=FONT, size=11)
            c.alignment = Alignment(wrap_text=True, vertical="top")
    h["A1"].font = Font(name=FONT, size=16, bold=True, color="1B2A4A")
    for r in range(3, h.max_row + 1):
        h.cell(r, 1).font = Font(name=FONT, size=11, bold=True)
    hdr = 16
    for c in (h.cell(hdr, 1), h.cell(hdr, 2)):
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = HEAD

    # --- Players ---
    ws = wb.create_sheet("Players")
    ws.append([c[0] for c in COLUMNS])
    for i, (name, width, _, _) in enumerate(COLUMNS, start=1):
        cell = ws.cell(1, i)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEAD
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = BOX
        ws.column_dimensions[cell.column_letter].width = width
        if name in NOTES:
            cell.comment = Comment(NOTES[name], "MindHunt")
    ws.row_dimensions[1].height = 34

    yn = lambda b: "Yes" if b else "No"
    for c in sorted(pack["candidates"], key=lambda c: (c["attrs"]["country"], c["name"])):
        a = c["attrs"]
        img = c.get("image") or {}
        ws.append([
            c["id"], c["name"], a["surname"], a["country"], a["role"], a["batting_hand"], a["bowling"],
            a.get("bowling_arm", ""), str(a.get("debut_decade", "N/A")), yn(a["captained_country"]), yn(a["odi_wc_winner"]),
            "Yes", "No", "", img.get("source_url", ""), img.get("license", ""), img.get("credit", ""),
        ])
    last_data = ws.max_row
    for _ in range(40):  # blank rows for new players
        ws.append([""] * len(COLUMNS))
    last = ws.max_row

    for r in range(2, last + 1):
        for i in range(1, len(COLUMNS) + 1):
            cell = ws.cell(r, i)
            cell.font = Font(name=FONT, size=10)
            cell.border = BOX
            cell.alignment = Alignment(vertical="center")
            if i == 1:
                cell.fill = LOCK
                cell.font = Font(name=FONT, size=9, color="6B7488")
            else:
                cell.fill = EDIT if r <= last_data else NEW

    # Dropdown lists on a reference sheet.
    ls = wb.create_sheet("Lists")
    col_of = {}
    for i, (key, vals) in enumerate(lists.items(), start=1):
        ls.cell(1, i, key).font = Font(name=FONT, bold=True)
        for j, v in enumerate(vals, start=2):
            ls.cell(j, i, v).font = Font(name=FONT)
        letter = ls.cell(1, i).column_letter
        col_of[key] = f"Lists!${letter}$2:${letter}${len(vals) + 1}"
    for i, (_, _, lst, strict) in enumerate(COLUMNS, start=1):
        if not lst:
            continue
        letter = ws.cell(1, i).column_letter
        dv = DataValidation(type="list", formula1=col_of[lst], allow_blank=True,
                            showErrorMessage=True, errorStyle="stop" if strict else "warning",
                            error="Pick a value from the list." if strict else "Not in the list. Keep it only if this is a new country.")
        ws.add_data_validation(dv)
        dv.add(f"{letter}2:{letter}{last}")

    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{ws.cell(1, len(COLUMNS)).column_letter}{last_data}"
    wb.active = 1
    wb.save(dst)
    print(f"{len(pack['candidates'])} players written to {dst}")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
