#!/usr/bin/env python3
"""Verification spreadsheets for any pack built by tools/build_packs.py.

  export:  python3 tools/pack_sheet.py export tools/packs/countries.py countries.xlsx
  import:  python3 tools/pack_sheet.py import tools/packs/countries.py edited.xlsx

Import checks every row, prints problems with row numbers and changes nothing
if there are any. Otherwise it rewrites the ROWS table in the spec file (the
source of truth) and stores image credits in <spec>.images.json; run
tools/build_packs.py afterwards. (Cricket has its own pair of scripts.)
"""
import json
import pathlib
import re
import runpy
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

FONT = "Arial"
HEAD = PatternFill("solid", fgColor="1B2A4A")
EDIT = PatternFill("solid", fgColor="FFF8DC")
LOCK = PatternFill("solid", fgColor="E6E9EE")
NEW = PatternFill("solid", fgColor="EAF5EE")
thin = Side(style="thin", color="C9D1DC")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
EXTRA = ["Keep in game", "Verified", "Source / notes", "Image page URL", "Image licence", "Image credit"]


def load_spec(path):
    mod = runpy.run_path(str(path))
    spec = mod["SPEC"]
    attrs = {a["key"]: a for a in spec["attributes"]}
    cols = [attrs[k] for k in spec["columns"][1:]]
    rows = []
    for line in mod["ROWS"].strip().splitlines():
        if line.strip() and not line.lstrip().startswith("#"):
            rows.append([c.strip() for c in line.split("|")])
    return spec, cols, rows


def to_cell(attr, raw):
    if raw in ("", "-"):
        return ""
    if attr["type"] == "bool":
        return "Yes" if raw == "Y" else "No"
    if attr["type"] == "set":
        return "; ".join(x.strip() for x in raw.split(";"))
    return raw


def export(spec_path, out):
    spec, cols, rows = load_spec(spec_path)
    wb = Workbook()
    h = wb.active
    h.title = "How to fill"
    h.column_dimensions["A"].width = 24
    h.column_dimensions["B"].width = 100
    info = [
        (f"MindHunt pack: {spec['title']}", ""),
        ("", ""),
        ("What this is", f"Every {spec.get('noun', 'candidate')} in this pack and the facts players can ask about. The game answers from this sheet, so a wrong cell means a wrong answer."),
        ("What to do", "1. Check each row in the Pack tab and fix anything wrong (yellow cells)."),
        ("", "2. Set Verified to Yes for rows you checked; put your source (a link) in Source / notes."),
        ("", "3. Add new rows in the green area at the bottom."),
        ("", "4. To remove a row, set Keep in game to No. Don't delete rows or change the header."),
        ("", "5. Leave a cell blank if the answer is genuinely unclear; the game then answers Unknown."),
        ("", "6. Optional: a Wikimedia Commons image page, its licence and credit."),
        ("Lists", "Columns with several values (like Continent or Field) take values separated by semicolons, e.g. Asia; Europe."),
        ("", ""),
        ("Column", "Meaning and allowed values"),
    ]
    for a in cols:
        allowed = ""
        if a["type"] in ("enum", "set"):
            allowed = " Allowed: " + ", ".join(a["values"]) + "."
        elif a["type"] == "bool":
            allowed = " Yes or No."
        elif a["type"] == "int":
            allowed = f" A number from {a['min']} to {a['max']}" + (f" in steps of {a['step']}." if a.get("step", 1) > 1 else ".")
        info.append((a["label"], (a.get("help", "") + allowed).strip()))
    for r in info:
        h.append(list(r))
    for row in h.iter_rows():
        for c in row:
            c.font = Font(name=FONT, size=11)
            c.alignment = Alignment(wrap_text=True, vertical="top")
    h["A1"].font = Font(name=FONT, size=16, bold=True, color="1B2A4A")
    for r in range(3, h.max_row + 1):
        h.cell(r, 1).font = Font(name=FONT, size=11, bold=True)
    for c in (h.cell(12, 1), h.cell(12, 2)):
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = HEAD

    ws = wb.create_sheet("Pack")
    headers = ["Name"] + [a["label"] for a in cols] + EXTRA
    ws.append(headers)
    widths = [26] + [max(12, min(28, len(a["label"]) + 4)) for a in cols] + [10, 10, 34, 30, 14, 20]
    for i, name in enumerate(headers, start=1):
        c = ws.cell(1, i)
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        c.border = BOX
        ws.column_dimensions[c.column_letter].width = widths[i - 1]
        if 1 < i <= len(cols) + 1 and cols[i - 2].get("help"):
            c.comment = Comment(cols[i - 2]["help"], "MindHunt")
    ws.row_dimensions[1].height = 40
    for r in rows:
        ws.append([r[0]] + [to_cell(a, raw) for a, raw in zip(cols, r[1:])] + ["Yes", "No", "", "", "", ""])
    last_data = ws.max_row
    for _ in range(30):
        ws.append([""] * len(headers))
    last = ws.max_row
    for r in range(2, last + 1):
        for i in range(1, len(headers) + 1):
            c = ws.cell(r, i)
            c.font = Font(name=FONT, size=10)
            c.border = BOX
            c.fill = EDIT if r <= last_data else NEW

    ls = wb.create_sheet("Lists")
    lists = {"yesno": ["Yes", "No"]}
    for a in cols:
        if a["type"] == "enum":
            lists[a["key"]] = a["values"]
        elif a["type"] == "int" and (a["max"] - a["min"]) // max(1, a.get("step", 1)) <= 120:
            lists[a["key"]] = [str(v) for v in range(a["min"], a["max"] + 1, max(1, a.get("step", 1)))]
    ref = {}
    for i, (k, vals) in enumerate(lists.items(), start=1):
        ls.cell(1, i, k).font = Font(name=FONT, bold=True)
        for j, v in enumerate(vals, start=2):
            ls.cell(j, i, v)
        L = ls.cell(1, i).column_letter
        ref[k] = f"Lists!${L}$2:${L}${len(vals) + 1}"

    def validate(col, key):
        L = ws.cell(1, col).column_letter
        dv = DataValidation(type="list", formula1=ref[key], allow_blank=True, showErrorMessage=True,
                            errorStyle="stop", error="Pick a value from the list.")
        ws.add_data_validation(dv)
        dv.add(f"{L}2:{L}{last}")

    for i, a in enumerate(cols, start=2):
        if a["type"] == "bool":
            validate(i, "yesno")
        elif a["key"] in lists:
            validate(i, a["key"])
    validate(len(cols) + 2, "yesno")
    validate(len(cols) + 3, "yesno")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{ws.cell(1, len(headers)).column_letter}{last_data}"
    wb.active = 1
    wb.save(out)
    print(f"{len(rows)} rows written to {out}")


def s(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def from_cell(a, v, err):
    if v == "":
        return "-"
    t = a["type"]
    if t == "bool":
        if v not in ("Yes", "No"):
            err(f"{a['label']} must be Yes or No")
        return "Y" if v == "Yes" else "N"
    if t == "int":
        try:
            n = int(float(v))
        except ValueError:
            err(f"{a['label']} '{v}' is not a number")
            return "-"
        step = max(1, a.get("step", 1))
        if not a["min"] <= n <= a["max"] or (n - a["min"]) % step:
            err(f"{a['label']} {n} out of range")
        return str(n)
    if t == "set":
        items = [x.strip() for x in re.split(r"[;,]", v) if x.strip()]
        for x in items:
            if x not in a["values"]:
                err(f"{a['label']} value '{x}' not allowed")
        return ";".join(items)
    if t == "enum" and v not in a["values"]:
        err(f"{a['label']} '{v}' not allowed")
    if "|" in v:
        err(f"{a['label']} may not contain |")
    return v


def do_import(spec_path, xlsx):
    spec, cols, _ = load_spec(spec_path)
    ws = load_workbook(xlsx, data_only=True)["Pack"]
    headers = ["Name"] + [a["label"] for a in cols] + EXTRA
    got = [s(c.value) for c in ws[1]][: len(headers)]
    if got != headers:
        sys.exit(f"Header changed. Expected: {headers}")
    errors, out, images, names, verified = [], [], [], set(), 0
    for r, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        v = [s(x) for x in row][: len(headers)] + [""] * len(headers)
        name = v[0]
        if not name:
            if any(v[1: len(cols) + 1]):
                errors.append(f"row {r}: has data but no Name")
            continue
        extra = dict(zip(EXTRA, v[len(cols) + 1: len(cols) + 1 + len(EXTRA)]))
        if extra["Keep in game"] == "No":
            continue
        err = lambda m, r=r, name=name: errors.append(f"row {r} ({name}): {m}")
        if "|" in name:
            err("name may not contain |")
        if name.lower() in names:
            err("duplicate name")
        names.add(name.lower())
        cells = [from_cell(a, val, err) for a, val in zip(cols, v[1: len(cols) + 1])]
        out.append([name] + cells)
        verified += extra["Verified"] == "Yes"
        if extra["Image page URL"]:
            if not extra["Image licence"]:
                err("image has no licence")
            images.append({"name": name, "source_url": extra["Image page URL"],
                           "license": extra["Image licence"], "credit": extra["Image credit"]})
    if errors:
        print(f"{len(errors)} problem(s); nothing changed:")
        print("\n".join(errors))
        sys.exit(1)
    widths = [max(len(row[i]) for row in out) for i in range(len(out[0]))]
    table = "\n".join(" | ".join(c.ljust(widths[i]) for i, c in enumerate(row)).rstrip() for row in out)
    text = pathlib.Path(spec_path).read_text()
    head, _, _ = text.partition('ROWS = """')
    new = head + 'ROWS = """\n' + table + '\n"""\n'
    if verified == len(out):
        new = new.replace('"as_of":', '"draft": False,\n    "as_of":', 1) if '"draft"' not in new else new
    pathlib.Path(spec_path).write_text(new)
    pathlib.Path(str(spec_path).replace(".py", ".images.json")).write_text(json.dumps(images, indent=2, ensure_ascii=False))
    print(f"{len(out)} rows, {verified} verified, {len(images)} images. Now run tools/build_packs.py (and bump the version if facts changed).")


if __name__ == "__main__":
    cmd, spec_path, file = sys.argv[1:4]
    export(spec_path, file) if cmd == "export" else do_import(spec_path, file)
