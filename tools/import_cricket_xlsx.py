#!/usr/bin/env python3
"""Turn the edited cricket sheet back into a pack.

Run: python3 tools/import_cricket_xlsx.py edited.xlsx server/packs/cricket-men-intl.v3.json OUT.json
Accepts the original sheet and data/cricket/mindhunt-cricket-players-verified.xlsx
(whose extra columns are ignored). ODI debut decade may be N/A (never played
an ODI); the game then answers Unknown.
Prints every problem with its row number and writes nothing if any are found.
Image columns are written to OUT.images.json for a separate download step.
"""
import hashlib
import json
import re
import sys

from openpyxl import load_workbook

HEADERS = ["ID", "Name", "Surname", "Country", "Role", "Batting hand", "Bowling", "Bowling arm",
           "Debut decade", "Captained country", "ODI World Cup winner", "Keep in game", "Verified",
           "Source / notes", "Image page URL", "Image licence", "Image credit"]


def s(v):
    return "" if v is None else str(v).strip()


def main(xlsx, base_pack, out):
    pack = json.load(open(base_pack))
    attrs = {a["key"]: a for a in pack["attributes"]}
    ws = load_workbook(xlsx, data_only=True)["Players"]
    head = [s(c.value) for c in ws[1]][: len(HEADERS)]
    head = ["Debut decade" if h == "ODI debut decade" else h for h in head]
    if head != HEADERS:
        sys.exit(f"Players header changed; expected {HEADERS}")

    errors, cands, images, seen_names, seen_ids = [], [], [], set(), set()
    countries = list(attrs["country"]["values"])
    verified = 0
    for r, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        v = dict(zip(HEADERS, (s(x) for x in row)))
        if not v["Name"]:
            if any(v[k] for k in HEADERS[2:11]):
                errors.append(f"row {r}: has data but no Name")
            continue
        if v["Keep in game"] == "No":
            continue
        err = lambda msg: errors.append(f"row {r} ({v['Name']}): {msg}")
        for key, field in [("role", "Role"), ("batting_hand", "Batting hand"), ("bowling", "Bowling")]:
            if v[field] not in attrs[key]["values"]:
                err(f"{field} '{v[field]}' not allowed")
        if not v["Country"]:
            err("Country missing")
        elif v["Country"] not in countries:
            countries.append(v["Country"])
        bowls = v["Bowling"] in ("Pace", "Spin")
        if bowls and v["Bowling arm"] not in ("Left", "Right"):
            err("Bowling arm must be Left or Right for a bowler")
        if not bowls and v["Bowling arm"]:
            err("Bowling arm must be blank for 'Not a regular bowler'")
        dec = None
        if v["Debut decade"].upper() not in ("N/A", "NA"):
            try:
                dec = int(float(v["Debut decade"]))
                if dec % 10 or not 1970 <= dec <= 2020:
                    raise ValueError
            except ValueError:
                err(f"ODI debut decade '{v['Debut decade']}' must be like 1990 (first ODI), or N/A")
        for f in ("Captained country", "ODI World Cup winner"):
            if v[f] not in ("Yes", "No"):
                err(f"{f} must be Yes or No")
        if not v["Surname"] or not re.search(r"[A-Za-z]", v["Surname"]):
            err("Surname missing")
        if v["Name"].lower() in seen_names:
            err("duplicate name")
        seen_names.add(v["Name"].lower())
        cid = v["ID"] or "c_" + hashlib.sha256(("mindhunt-cricket:" + v["Name"]).encode()).hexdigest()[:6]
        if cid in seen_ids:
            err(f"duplicate ID {cid}")
        seen_ids.add(cid)
        if v["Verified"] == "Yes":
            verified += 1
        a = {"country": v["Country"], "role": v["Role"], "batting_hand": v["Batting hand"], "bowling": v["Bowling"],
             "captained_country": v["Captained country"] == "Yes",
             "odi_wc_winner": v["ODI World Cup winner"] == "Yes", "surname": v["Surname"]}
        if dec is not None:
            a["debut_decade"] = dec
        if bowls:
            a["bowling_arm"] = v["Bowling arm"]
        cands.append({"id": cid, "name": v["Name"], "attrs": a})
        if v["Image page URL"]:
            if not v["Image licence"]:
                err("image has no licence")
            images.append({"id": cid, "name": v["Name"], "source_url": v["Image page URL"],
                           "license": v["Image licence"], "credit": v["Image credit"]})

    if errors:
        print(f"{len(errors)} problem(s); nothing written:")
        print("\n".join(errors))
        sys.exit(1)
    attrs["country"]["values"] = countries
    attrs["debut_decade"]["min"] = min([1970] + [c["attrs"]["debut_decade"] for c in cands if "debut_decade" in c["attrs"]])
    old = {c["id"]: c for c in pack["candidates"]}
    order = {c["id"]: i for i, c in enumerate(pack["candidates"])}  # keep the pack's order; new players last
    cands.sort(key=lambda c: order.get(c["id"], len(order)))
    changed = cands != [old.get(c["id"]) for c in cands] or len(cands) != len(old)
    if changed:
        pack["version"] += 1
    pack["candidates"] = cands
    pack["draft"] = verified < len(cands)
    json.dump(pack, open(out, "w"), indent=2, ensure_ascii=False)
    json.dump(images, open(out.replace(".json", "") + ".images.json", "w"), indent=2, ensure_ascii=False)
    print(f"{len(cands)} players, {verified} verified, {len(images)} images, version {pack['version']}"
          f"{'' if changed else ' (no data changes)'}")


if __name__ == "__main__":
    main(*sys.argv[1:4])
