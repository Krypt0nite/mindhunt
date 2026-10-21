#!/usr/bin/env python3
"""Check the cricket pack against its evidence and build the verification deliverables.

Run: python3 tools/verify_cricket.py [original.xlsx]

Fails (exit 1, nothing written) if any pack value disagrees with the evidence
in tools/cricket_evidence.py, an ID changed, or an image entry breaks the
attribution rules. Otherwise writes, under data/cricket/:

  mindhunt-cricket-players-verified.xlsx  Players, Audit, Image attribution, Methodology, Lists
  cricket-audit.csv / cricket-audit.json   one record per player and checked field
  cricket-players.json                     flat app-ready dataset (the game itself loads
                                           server/packs/cricket-men-intl.v3.json)
  cricket-image-attribution.csv            the image manifest as a table

The original workbook is read, never modified; its values are the audit's
"original value" column.
"""
import csv
import json
import pathlib
import re
import runpy
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = pathlib.Path(__file__).resolve().parent.parent
OUT = ROOT / "data" / "cricket"
PACK = ROOT / "server" / "packs" / "cricket-men-intl.v3.json"
MANIFEST = ROOT / "tools" / "portraits" / "cricket-manifest.json"
ASSETS = ROOT / "assets"
EV = runpy.run_path(str(ROOT / "tools" / "cricket_evidence.py"))
P, INFOBOX, WIKI = EV["P"], EV["INFOBOX"], EV["WIKI"]

HEADERS = ["ID", "Name", "Surname", "Country", "Role", "Batting hand", "Bowling", "Bowling arm",
           "ODI debut decade", "Captained country", "ODI World Cup winner", "Keep in game", "Verified",
           "Source / notes", "Image page URL", "Image licence", "Image credit",
           "Data verification", "Image verification", "First ODI", "First ODI team", "Captaincy formats",
           "World Cup titles", "Review note"]
FIELDS = ["Name", "Surname", "Country", "Role", "Batting hand", "Bowling", "Bowling arm",
          "ODI debut decade", "Captained country", "ODI World Cup winner"]

FONT = "Arial"
HEAD = PatternFill("solid", fgColor="1B2A4A")
OK = PatternFill("solid", fgColor="EAF5EE")
WARN = PatternFill("solid", fgColor="FFF4CC")
BAD = PatternFill("solid", fgColor="FBE3E1")
thin = Side(style="thin", color="C9D1DC")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
FILL = {"Verified": OK, "Corrected": OK, "Needs review": WARN, "Unverified": BAD,
        "Missing": BAD, "Rejected": BAD}


def wurl(title):
    return WIKI + title


def derived_bowling(style):
    """(class, arm) from an infobox bowling style, for players who bowl."""
    s = style.lower()
    kind = "Pace" if re.search(r"fast|medium|seam", s) else "Spin" if re.search(r"break|orthodox|spin|googly", s) else None
    arm = "Left" if "left-arm" in s else "Right" if kind else None
    return kind, arm


def pack_rows():
    pack = json.loads(PACK.read_text())
    out = {}
    for c in pack["candidates"]:
        a = c["attrs"]
        yn = lambda b: "Yes" if b else "No"
        out[c["id"]] = {
            "ID": c["id"], "Name": c["name"], "Surname": a["surname"], "Country": a["country"], "Role": a["role"],
            "Batting hand": a["batting_hand"], "Bowling": a["bowling"], "Bowling arm": a.get("bowling_arm", ""),
            "ODI debut decade": str(a["debut_decade"]) if "debut_decade" in a else "N/A",
            "Captained country": yn(a["captained_country"]), "ODI World Cup winner": yn(a["odi_wc_winner"]),
            "_image": c.get("image"),
        }
    return pack, out


def original_rows(path):
    ws = load_workbook(path, data_only=True)["Players"]
    rows = list(ws.iter_rows(values_only=True))
    head = [str(h) for h in rows[0]]
    out = []
    for r in rows[1:]:
        if not r[1]:
            continue
        d = {h: ("" if v is None else str(v)) for h, v in zip(head, r)}
        d["ODI debut decade"] = d.get("Debut decade", "")
        out.append(d)
    return out


def audit_player(row, orig, ev, errors):
    """Audit records for one player; appends to errors where pack and evidence disagree."""
    name = row["Name"]
    art = wurl(ev["wiki"])
    country = row["Country"]
    capt_list = wurl(EV["CAPTAIN_LISTS"][country])
    odi_list = wurl(EV["ODI_LISTS"][country])
    bat, style = INFOBOX[name]
    recs = []

    def rec(field, status, src, src2, why, expected=None):
        final = row[field]
        if expected is not None and expected != final:
            errors.append(f"{name}: {field} is {final!r} but the evidence says {expected!r}")
        original = orig.get(field, "")
        if status == "Verified" and original != final:
            status = "Corrected"
        recs.append({"Player ID": row["ID"], "Player name": name, "Field": field, "Original value": original,
                     "Final value": final, "Source URL": src, "Secondary source URL": src2,
                     "Status": status, "Explanation": why})

    old = ev.get("old_name")
    rec("Name", "Verified", art, "",
        f"Spelling corrected from '{old}' to the article's spelling; ID kept from the old name." if old else "Matches the Wikipedia article title.")
    rec("Surname", "Verified", art, "", "Follows the corrected name." if old else "From the name.")
    rec("Country", "Verified", odi_list, art,
        f"Listed among {country} ODI players." + (f" First ODI was for {ev['odi_team']}." if ev.get("odi_team") else "")
        if ev["odi"] else f"Represented {country} in Tests (no ODIs).")

    role_note = ev.get("wiki_role")
    decided = EV.get("ROLE_DECISIONS", {}).get(name)
    if role_note and decided:
        rec("Role", "Verified", art, "",
            f"Infobox calls him '{role_note}'; the pack owner chose '{decided}' on {EV['ROLE_DECIDED_ON']} (the pack has 4 classes).",
            expected=decided)
        role_note = None
    else:
        rec("Role", "Needs review" if role_note else "Verified", art, "",
            f"Infobox calls him '{role_note}'; pack keeps '{row['Role']}' (the pack's 4 classes). Needs a decision." if role_note
            else "Consistent with the infobox playing role.")

    rec("Batting hand", "Verified", art, "", f"Infobox: {'right' if bat == 'R' else 'left'}-handed.",
        expected={"R": "Right", "L": "Left"}[bat])

    if row["Role"] in ("Batter", "Wicketkeeper"):
        why = "Workbook rule: specialist batters and keepers are 'Not a regular bowler'" + (f" (infobox lists occasional {style.lower()})." if style else ".")
        rec("Bowling", "Needs review" if role_note else "Verified", art, "", why + (" Depends on the role review." if role_note else ""),
            expected="Not a regular bowler")
        rec("Bowling arm", "Verified", art, "", "Blank because not a regular bowler.", expected="")
    else:
        kind, arm = derived_bowling(style)
        rec("Bowling", "Verified", art, "", f"Infobox: {style}.", expected=kind)
        rec("Bowling arm", "Verified", art, "", f"Infobox: {style}.", expected=arm)

    if ev["odi"] is None:
        rec("ODI debut decade", "Verified", art, odi_list,
            "Never played an ODI (career 1928-48; ODIs began in 1971), so N/A and the game answers Unknown.", expected="N/A")
    else:
        y = int(ev["odi"][:4])
        why = f"First ODI {ev['odi']}" + (f" for {ev['odi_team']}" if ev.get("odi_team") else "") + "."
        if ev.get("note") and ("Test debut" in ev["note"] or "first for" in ev["note"]):
            why += " " + ev["note"]
        rec("ODI debut decade", "Verified", art, odi_list, why, expected=str(y // 10 * 10))

    fm = ", ".join(ev["capt"])
    rec("Captained country", "Verified", capt_list, art,
        f"Listed as {country} captain in: {fm}." if ev["capt"] else f"Not in the {country} captains list (Test, ODI or T20I).",
        expected="Yes" if ev["capt"] else "No")

    if ev["wc"]:
        srcs = EV["WC_SOURCES"][ev["wc"][-1]]
        why = "Member of the winning squad in " + ", ".join(map(str, ev["wc"])) + "."
        if ev.get("wc_unchecked"):
            why += f" Also believed to have won {', '.join(map(str, ev['wc_unchecked']))}, not re-checked here."
        if "Warne" in name or "Waqar" in name:
            why += " " + ev["note"]
        rec("ODI World Cup winner", "Verified", srcs[0], srcs[1] if len(srcs) > 1 else "", why, expected="Yes")
    else:
        won = [y for y, c in EV["WINNERS"].items() if c == country]
        why = (f"{country} won in {', '.join(map(str, won))}; not in those squads." if won
               else f"{country} has never won the men's ODI World Cup.")
        if "Waqar" in name:
            why += " " + ev["note"]
        src = EV["WC_SOURCES"][won[-1]][0] if won else EV["WINNERS_URL"]
        rec("ODI World Cup winner", "Verified", src, EV["WINNERS_URL"] if won else "", why, expected="No")
    return recs


def check_images(manifest, pack, errors):
    seen = {}
    pack_imgs = {c["id"]: c.get("image") for c in pack["candidates"]}
    for e in manifest:
        if e["image_status"] not in ("Verified", "Missing", "Rejected", "Needs review"):
            errors.append(f"{e['name']}: bad image status {e['image_status']}")
        if e["image_status"] == "Verified":
            for k in ("file_page", "credit", "license", "asset", "source_url"):
                if not e.get(k):
                    errors.append(f"{e['name']}: verified image missing {k}")
            if not e.get("identity_confirmed") or e.get("licence_status") != "Verified":
                errors.append(f"{e['name']}: verified image without confirmed identity and licence")
            f = ASSETS / e.get("asset", "")
            try:
                from PIL import Image
                with Image.open(f) as im:
                    im.verify()
            except Exception as ex:  # noqa: BLE001
                errors.append(f"{e['name']}: asset {f} does not decode: {ex}")
            if not pack_imgs.get(e["id"]):
                errors.append(f"{e['name']}: verified image not in the pack; rerun gen_cricket_pack.py")
        elif pack_imgs.get(e["id"]):
            errors.append(f"{e['name']}: pack has an image that is not verified")
        for k in ("sha1", "file_page"):
            v = e.get(k)
            if v and e["image_status"] in ("Verified", "Needs review"):
                if v in seen and seen[v] != e["id"]:
                    errors.append(f"{e['name']}: same image as {seen[v]}")
                seen[v] = e["id"]


def style_sheet(ws, widths, status_cols=()):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    for c in ws[1]:
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = BOX
    ws.row_dimensions[1].height = 32
    idx = [i for i, c in enumerate(ws[1], start=1) if c.value in status_cols]
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name=FONT, size=10)
            c.border = BOX
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for i in idx:
            c = row[i - 1]
            if c.value in FILL:
                c.fill = FILL[c.value]
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions


def main(original):
    pack, rows = pack_rows()
    orig = original_rows(original)
    errors = []
    if len(rows) != 72 or len(orig) != 72:
        errors.append(f"expected 72 players, pack has {len(rows)}, original sheet {len(orig)}")
    if set(rows) != {o["ID"] for o in orig}:
        errors.append("player IDs differ from the original workbook: " + str(set(rows) ^ {o["ID"] for o in orig}))
    if len(set(rows)) != len(rows):
        errors.append("duplicate IDs")
    if sorted(P) != sorted(r["Name"] for r in rows.values()):
        errors.append("evidence and pack rosters differ")

    manifest = json.loads(MANIFEST.read_text())
    mby = {e["id"]: e for e in manifest}
    check_images(manifest, pack, errors)

    audit, players, flat = [], [], []
    for o in orig:  # original workbook order
        r = rows.get(o["ID"])
        if r is None:
            continue
        ev = P[r["Name"]]
        recs = audit_player(r, o, ev, errors)
        audit += recs
        statuses = {x["Status"] for x in recs}
        data_status = "Needs review" if "Needs review" in statuses else "Unverified" if "Unverified" in statuses else "Verified"
        for x in recs:
            if not x["Source URL"]:
                errors.append(f"{r['Name']}: {x['Field']} has no source URL")
        m = mby.get(r["ID"], {})
        img_ok = m.get("image_status") == "Verified"
        review = "; ".join(x["Explanation"] for x in recs if x["Status"] == "Needs review" and x["Field"] == "Role")
        changed = [x for x in recs if x["Status"] == "Corrected"]
        notes = (" Corrected: " + "; ".join(f"{x['Field']} {x['Original value'] or 'blank'} -> {x['Final value']}" for x in changed) + ".") if changed else ""
        players.append({
            **{k: r[k] for k in HEADERS[:11]},
            "Keep in game": "Yes", "Verified": "Yes" if data_status == "Verified" else "No",
            "Source / notes": f"Checked {EV['ACCESSED']}; per-field sources on the Audit sheet.{notes}",
            "Image page URL": m.get("file_page", "") if img_ok else "",
            "Image licence": m.get("license", "") if img_ok else "",
            "Image credit": m.get("credit", "") if img_ok else "",
            "Data verification": data_status, "Image verification": m.get("image_status", "Missing"),
            "First ODI": ev["odi"] or "Never played an ODI", "First ODI team": ev.get("odi_team", r["Country"] if ev["odi"] else ""),
            "Captaincy formats": ", ".join(ev["capt"]) or "None",
            "World Cup titles": ", ".join(map(str, ev["wc"])) + (f" (+{', '.join(map(str, ev['wc_unchecked']))} not re-checked)" if ev.get("wc_unchecked") else "") or "None",
            "Review note": review,
        })
        flat.append({"id": r["ID"], "name": r["Name"], "surname": r["Surname"], "country": r["Country"], "role": r["Role"],
                     "batting_hand": r["Batting hand"], "bowling": r["Bowling"], "bowling_arm": r["Bowling arm"] or None,
                     "odi_debut_decade": None if r["ODI debut decade"] == "N/A" else int(r["ODI debut decade"]),
                     "odi_debut_date": ev["odi"], "captained_country": r["Captained country"] == "Yes",
                     "captaincy_formats": ev["capt"], "odi_wc_winner": r["ODI World Cup winner"] == "Yes",
                     "odi_wc_titles": ev["wc"], "data_verification": data_status,
                     "image_verification": m.get("image_status", "Missing"), "image": r["_image"]})

    if errors:
        print(f"{len(errors)} problem(s); nothing written:")
        print("\n".join(errors))
        sys.exit(1)

    OUT.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Players"
    ws.append(HEADERS)
    for p in players:
        ws.append([p[h] for h in HEADERS])
    style_sheet(ws, [11, 22, 14, 13, 13, 9, 15, 9, 10, 10, 10, 8, 8, 40, 26, 12, 16, 13, 13, 12, 12, 14, 16, 50],
                ("Data verification", "Image verification"))
    lists = {"Role": ["Batter", "Bowler", "All-rounder", "Wicketkeeper"], "Batting hand": ["Left", "Right"],
             "Bowling": ["Pace", "Spin", "Not a regular bowler"], "Bowling arm": ["Left", "Right"],
             "ODI debut decade": ["N/A", "1970", "1980", "1990", "2000", "2010", "2020"],
             "Captained country": ["Yes", "No"], "ODI World Cup winner": ["Yes", "No"], "Keep in game": ["Yes", "No"],
             "Verified": ["Yes", "No"], "Data verification": ["Verified", "Needs review", "Unverified"],
             "Image verification": ["Verified", "Missing", "Rejected", "Needs review"]}
    ls = wb.create_sheet("Lists")
    for i, (k, vals) in enumerate(lists.items(), start=1):
        ls.cell(1, i, k).font = Font(name=FONT, bold=True)
        for j, v in enumerate(vals, start=2):
            ls.cell(j, i, v)
        L = ls.cell(1, i).column_letter
        dv = DataValidation(type="list", formula1=f"Lists!${L}$2:${L}${len(vals) + 1}", allow_blank=True)
        ws.add_data_validation(dv)
        col = ws.cell(1, HEADERS.index(k) + 1).column_letter
        dv.add(f"{col}2:{col}{len(players) + 31}")

    au = wb.create_sheet("Audit", 1)
    acols = ["Player ID", "Player name", "Field", "Original value", "Final value", "Source URL",
             "Secondary source URL", "Status", "Explanation"]
    au.append(acols)
    for a in audit:
        au.append([a[c] for c in acols])
    style_sheet(au, [11, 22, 18, 14, 14, 48, 48, 13, 70], ("Status",))

    im = wb.create_sheet("Image attribution", 2)
    icols = ["Player ID", "Player name", "Image verification", "File page", "Direct image URL", "Creator",
             "Licence", "Licence URL", "Attribution line", "Identity verification", "Licence verification",
             "Download / decode", "Candidates examined", "Notes"]
    im.append(icols)
    for e in sorted(manifest, key=lambda e: [o["ID"] for o in orig].index(e["id"])):
        im.append([e["id"], e["name"], e["image_status"], e["file_page"], e["file_url"], e["credit"], e["license"],
                   e["license_url"], e["attribution_line"], "Confirmed" if e["identity_confirmed"] else "Not confirmed",
                   e["licence_status"], e["download_status"],
                   "\n".join(f"{c['file']}: {c['result']}" for c in e["candidates_examined"]), e["reason"]])
    style_sheet(im, [11, 22, 13, 40, 30, 20, 14, 30, 40, 14, 13, 18, 60, 60], ("Image verification",))

    me = wb.create_sheet("Methodology", 3)
    for row in methodology(players, audit, manifest):
        me.append(row)
    me.column_dimensions["A"].width = 30
    me.column_dimensions["B"].width = 120
    for row in me.iter_rows():
        for c in row:
            c.font = Font(name=FONT, size=10, bold=c.column == 1)
            c.alignment = Alignment(wrap_text=True, vertical="top")
    me["A1"].font = Font(name=FONT, size=14, bold=True, color="1B2A4A")

    xlsx = OUT / "mindhunt-cricket-players-verified.xlsx"
    wb.save(xlsx)
    with open(OUT / "cricket-audit.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=acols)
        w.writeheader()
        w.writerows(audit)
    (OUT / "cricket-audit.json").write_text(json.dumps(audit, indent=1, ensure_ascii=False) + "\n")
    (OUT / "cricket-players.json").write_text(json.dumps(
        {"pack": "cricket-men-intl", "version": pack["version"], "checked": EV["ACCESSED"], "players": flat},
        indent=1, ensure_ascii=False) + "\n")
    with open(OUT / "cricket-image-attribution.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(icols)
        for r in im.iter_rows(min_row=2, values_only=True):
            w.writerow(r)

    st = lambda k, v: sum(1 for x in audit if x[k] == v)
    print(f"72 players: {sum(p['Data verification'] == 'Verified' for p in players)} fully verified, "
          f"{sum(p['Data verification'] == 'Needs review' for p in players)} need review")
    print(f"{len(audit)} facts: {st('Status', 'Verified')} verified, {st('Status', 'Corrected')} corrected, "
          f"{st('Status', 'Needs review')} need review")
    img = {}
    for e in manifest:
        img[e["image_status"]] = img.get(e["image_status"], 0) + 1
    print("images:", ", ".join(f"{k} {v}" for k, v in sorted(img.items())))
    print(f"wrote {xlsx.relative_to(ROOT)} and data/cricket/*.csv|json")


def methodology(players, audit, manifest):
    corr = [a for a in audit if a["Status"] == "Corrected"]
    rev = [a for a in audit if a["Status"] == "Needs review"]
    examined = sum(1 for e in manifest for c in e["candidates_examined"] if not c["file"].startswith("("))
    return [
        ("MindHunt cricket: verification method", ""),
        ("Checked", f"{EV['ACCESSED']}. 72 players, 10 fields each ({len(audit)} audit records)."),
        ("", ""),
        ("Definitions", ""),
        ("ODI debut decade", "Decade of the player's first one-day international, for any country (Eoin Morgan: Ireland 2006; Dale Steyn: Africa XI 2005). Test and T20I debuts are ignored. Don Bradman never played an ODI, so the value is N/A and the game answers Unknown to that question. The pack's range is now 1970-2020."),
        ("Captained country", "Yes if the player is listed as an official captain of his country in at least one format (Test, ODI or T20I), including stand-in matches. The formats are in the Players sheet."),
        ("ODI World Cup winner", "Yes if he was in the squad of a team that won the men's 50-over World Cup (1975-2023). A player who withdrew before the tournament does not count (Shane Warne, 2003)."),
        ("Role", "One of four classes: Batter, Bowler, All-rounder, Wicketkeeper (keepers who mainly bat). Batters and keepers are 'Not a regular bowler' whatever they bowled occasionally (workbook rule)."),
        ("Bowling / arm", "Pace = fast, fast-medium, medium or seam; Spin = off break, leg break, orthodox or googly. Arm from the infobox ('left-arm', 'slow left-arm orthodox')."),
        ("Image verification", "Verified only when the Commons file page was opened, the licence is on the allow-list (CC0, public domain, CC BY, CC BY-SA), the author is recorded, the file was downloaded and decoded, and a person confirmed the photo shows the player. Missing = no usable image examined. Rejected = every candidate failed a rule. Needs review = a lead exists but a check is outstanding."),
        ("Statuses", "Verified = matches the cited source. Corrected = changed from the original workbook to match the source (original kept in the Audit sheet). Needs review = a judgement the source does not settle. Unverified = no source reached (none in this release)."),
        ("", ""),
        ("Sources used", "Official: BCCI 2011 squad announcement; ICC 2019 and 2023 squad articles. Secondary: English Wikipedia player articles (infobox and text), the country lists of captains and ODI cricketers, the World Cup squad pages and the list of World Cup finals. Each audit row cites the page the fact came from; a secondary URL is given where a second page agrees."),
        ("Source limits", "The brief asks for ICC/ESPNcricinfo first. From the build environment, ESPNcricinfo player profiles, Wikidata and Wikimedia Commons could not be opened (fetch errors / blocked by the network proxy), so most facts rest on Wikipedia. No disagreement was found where an official page was also reachable (2011, 2019, 2023 squads). Wikipedia article URLs are recorded in canonical form."),
        ("Not re-checked", "The 2015 Australian squad page did not render. Smith, Cummins and Starc are Yes on the strength of the 2023 squad; their 2015 titles are noted but not counted as checked."),
        ("", ""),
        ("Corrections", "; ".join(f"{a['Player name']}: {a['Field']} {a['Original value'] or 'blank'} -> {a['Final value']}" for a in corr)),
        ("Needs review", "; ".join(f"{a['Player name']} ({a['Field']}): {a['Explanation']}" for a in rev)),
        ("Images", f"{sum(e['image_status'] == 'Verified' for e in manifest)} verified. {examined} candidate file pages were tried (the leads in the brief) and none could be opened, so no licence or author was confirmed and none was downloaded. The earlier note 'CC BY 2.0, Jim Dowling' was not copied to any row. Run tools/fetch_portraits.py on a machine that can reach Commons, review the pictures, then confirm them."),
        ("Rebuild", "tools/cricket_evidence.py holds the evidence; tools/gen_cricket_pack.py the pack; python3 tools/verify_cricket.py rebuilds this workbook and fails if pack and evidence disagree."),
    ]


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else str(ROOT / "data" / "cricket" / "original" / "mindhunt-cricket-players.xlsx"))
