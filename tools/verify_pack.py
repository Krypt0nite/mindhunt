#!/usr/bin/env python3
"""Check a topic pack against its evidence and build its verification files.

  python3 tools/verify_pack.py countries [--baseline d569dca]

Reads tools/packs/<name>.py (the pack) and tools/evidence/<name>.py:
  EXPECT  {attr: {item name: value in pack notation}}  ("Y"/"N" for yes/no,
          "-" for Unknown, "A;B" for lists); an item missing from EXPECT[attr]
          means that fact was not checked
  SOURCES {attr: [url, ...]}, optional SOURCE_TIER {attr: text}
  RULES   {attr: how the definition was applied}
  NOTES   {(item, attr): explanation}, optional
  CHECKED date

Fails (exit 1, nothing written) if any checked fact disagrees with the pack.
Otherwise writes data/<pack id>/: <pack id>-verified.xlsx (Pack, Audit,
Methodology), audit.csv and audit.json. The audit's "original value" is the
pack as of --baseline (a git revision; default: the commit that added the pack).
"""
import argparse
import csv
import json
import pathlib
import runpy
import subprocess
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ROOT = pathlib.Path(__file__).resolve().parent.parent
HEAD = PatternFill("solid", fgColor="1B2A4A")
FILL = {"Verified": PatternFill("solid", fgColor="EAF5EE"), "Corrected": PatternFill("solid", fgColor="EAF5EE"),
        "Needs review": PatternFill("solid", fgColor="FFF4CC"), "Unverified": PatternFill("solid", fgColor="FBE3E1")}
thin = Side(style="thin", color="C9D1DC")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)


def rows_of(text):
    mod = {}
    exec(compile(text, "spec", "exec"), mod)
    spec = mod["SPEC"]
    out = {}
    for line in mod["ROWS"].strip().splitlines():
        if line.strip() and not line.lstrip().startswith("#"):
            cells = [c.strip() for c in line.split("|")]
            out[cells[0]] = dict(zip(spec["columns"][1:], cells[1:]))
    return spec, out


def show(attr, raw):
    if raw in ("", "-", None):
        return "Unknown"
    if attr["type"] == "bool":
        return "Yes" if raw == "Y" else "No"
    if attr["type"] == "set":
        return "; ".join(raw.split(";"))
    return raw


def same(attr, a, b):
    if attr["type"] == "set":
        return sorted(a.split(";")) == sorted(b.split(";"))
    return a == b


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("name")
    ap.add_argument("--baseline", default="")
    args = ap.parse_args()
    spec_path = ROOT / "tools" / "packs" / f"{args.name}.py"
    ev = runpy.run_path(str(ROOT / "tools" / "evidence" / f"{args.name}.py"))
    spec, rows = rows_of(spec_path.read_text())
    attrs = {a["key"]: a for a in spec["attributes"]}
    rel = spec_path.relative_to(ROOT).as_posix()
    rev = args.baseline or subprocess.run(["git", "log", "--diff-filter=A", "--format=%h", "--", rel], cwd=ROOT,
                                          capture_output=True, text=True).stdout.split()[-1]
    _, base = rows_of(subprocess.run(["git", "show", f"{rev}:{rel}"], cwd=ROOT, capture_output=True, text=True, check=True).stdout)

    errors, audit = [], []
    expect, sources, notes = ev["EXPECT"], ev.get("SOURCES", {}), ev.get("NOTES", {})
    for name, vals in rows.items():
        for key in spec["columns"][1:]:
            a = attrs[key]
            final = vals.get(key, "-")
            orig = base.get(name, {}).get(key, "")
            exp = expect.get(key, {}).get(name)
            src = sources.get(key, [""])
            if key in ev.get("PER_ITEM", ()):  # the item's own article first, then any list page
                src = [ev["ITEM_URL"][name]] + [u for u in sources.get(key, []) if u]
            if exp is None:
                status, why = "Unverified", "Not checked against a source yet."
            elif not same(a, final, exp):
                errors.append(f"{name}: {key} is {final!r}, evidence says {exp!r}")
                continue
            else:
                status = "Verified" if same(a, orig, final) or name not in base else "Corrected"
                why = notes.get((name, key), "")
                if status == "Corrected" and not why:
                    why = f"Changed from {show(a, orig)} to match the source."
                if not why:
                    why = "Matches the source."
            audit.append({"Item": name, "Field": a["label"], "Original value": show(a, orig) if name in base else "(new)",
                          "Final value": show(a, final), "Source URL": src[0], "Secondary source URL": src[1] if len(src) > 1 else "",
                          "Source tier": ev.get("SOURCE_TIER", {}).get(key, "Secondary (Wikipedia)") if exp is not None else "",
                          "Status": status, "Explanation": why})
    for key, m in expect.items():
        for name in m:
            if name not in rows:
                errors.append(f"evidence names {name!r} ({key}) but the pack has no such item")
    if errors:
        print(f"{len(errors)} problem(s); nothing written:")
        print("\n".join(errors))
        sys.exit(1)

    out = ROOT / "data" / spec["id"]
    out.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Pack"
    cols = [attrs[k] for k in spec["columns"][1:]]
    ws.append(["Name"] + [a["label"] for a in cols] + ["Data verification"])
    for name, vals in rows.items():
        sts = {x["Status"] for x in audit if x["Item"] == name}
        st = "Unverified" if "Unverified" in sts else "Needs review" if "Needs review" in sts else "Verified"
        ws.append([name] + [show(a, vals.get(a["key"], "-")) for a in cols] + [st])
    au = wb.create_sheet("Audit")
    acols = list(audit[0].keys())
    au.append(acols)
    for x in audit:
        au.append([x[c] for c in acols])
    me = wb.create_sheet("Methodology")
    me.append([f"{spec['title']}: verification", ""])
    me.append(["Checked", f"{ev['CHECKED']}. {len(rows)} items × {len(cols)} fields = {len(audit)} facts."])
    me.append(["", ""])
    for a in cols:
        me.append([a["label"], (a.get("help", "") + " " + ev.get("RULES", {}).get(a["key"], "")).strip()])
    me.append(["", ""])
    corr = [x for x in audit if x["Status"] == "Corrected"]
    me.append(["Corrections", "; ".join(f"{x['Item']}: {x['Field']} {x['Original value']} -> {x['Final value']}" for x in corr) or "None"])
    me.append(["Unknown answers", "Blank cells are answered Unknown in the game and never eliminate anyone."])
    for sheet, widths, status_col in ((ws, [22] + [16] * len(cols) + [14], "Data verification"),
                                      (au, [22, 22, 14, 14, 48, 40, 22, 12, 60], "Status"), (me, [26, 110], None)):
        for i, w in enumerate(widths, start=1):
            sheet.column_dimensions[sheet.cell(1, i).column_letter].width = w
        for c in sheet[1]:
            c.font = Font(name="Arial", bold=True, color="FFFFFF")
            c.fill = HEAD
        idx = [c.column for c in sheet[1] if c.value == status_col]
        for row in sheet.iter_rows(min_row=2):
            for c in row:
                c.font = Font(name="Arial", size=10, bold=sheet is me and c.column == 1)
                c.alignment = Alignment(wrap_text=True, vertical="top")
                c.border = BOX
            for i in idx:
                if row[i - 1].value in FILL:
                    row[i - 1].fill = FILL[row[i - 1].value]
        if sheet is not me:
            sheet.freeze_panes = "B2"
            sheet.auto_filter.ref = sheet.dimensions
    xlsx = out / f"{spec['id']}-verified.xlsx"
    wb.save(xlsx)
    with open(out / "audit.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=acols)
        w.writeheader()
        w.writerows(audit)
    (out / "audit.json").write_text(json.dumps(audit, indent=1, ensure_ascii=False) + "\n")
    count = lambda s: sum(x["Status"] == s for x in audit)
    print(f"{spec['id']}: {len(rows)} items, {len(audit)} facts: {count('Verified')} verified, {count('Corrected')} corrected, "
          f"{count('Needs review')} need review, {count('Unverified')} not checked")
    print(f"wrote {xlsx.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
